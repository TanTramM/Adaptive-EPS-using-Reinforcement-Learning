# -*- coding: utf-8 -*-
"""Build the English progress report on the common parameter-selection method (Opt controllers) for the supervisor.

Reads only existing result files (no simulation):
  Model/data/{map,pid,mappid,mapopt,pidopt,mappidopt}.json         designs and CALIB scores
  Result/Controllers/<Ctrl>/Cases/<Ctrl>_metrics_s9999[5-9].csv     DREWS on the five scoring seeds
  Result/Controllers/<Ctrl>/Cases/<Ctrl>_SLIP_sig_s99999.csv        time responses (one seed)
  Result/Controllers/{MapOpt,PIDOpt,MapPIDOpt}/Calib/*_CALIB_*.csv  search evaluations (Pareto) and the multi-seed choice
Writes figures to Result/Report/Opt_*.png, tables to Result/Report/Opt_*.csv and the report to Documents/Report/Opt_method_report.docx.
Usage: python -I -X utf8 tools/build_report_opt.py
"""
import csv, glob, json, math, os, statistics as st
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RES = os.path.join(ROOT, 'Result')
FIG = os.path.join(RES, 'Report'); os.makedirs(FIG, exist_ok=True)
OUT = os.path.join(ROOT, 'Documents', 'Report', 'Opt_method_report.docx')
SEEDS = [99999, 99998, 99997, 99996, 99995]
CASES = ['DRY', 'RAMP', 'SLIP', 'WEAVE', 'EDGE']
CTRL = ['MapOpt', 'PIDOpt', 'MapPIDOpt']   # results only for the -Opt controllers (user 2026-10-17); the grid designs appear in the method section only
LABEL = {'MapOpt': 'Map', 'PIDOpt': 'PID', 'MapPIDOpt': 'MapPID'}
COLOR = {'MapOpt': '#2171b5', 'PIDOpt': '#d94801', 'MapPIDOpt': '#238b45'}


def jload(name):
    with open(os.path.join(ROOT, 'Model', 'data', name + '.json'), encoding='utf-8') as f:
        return json.load(f)['design']


def metrics(ctrl, seed):
    p = os.path.join(RES, 'Controllers', ctrl, 'Cases', '%s_metrics_s%05d.csv' % (ctrl, seed))
    if not os.path.exists(p):
        return None
    D = {}
    for r in csv.DictReader(open(p, encoding='utf-8')):
        c = r['Case']
        if r['Window'] == 'whole case':
            D.setdefault(c, {}).update(rms=float(r['RMS_eT_Nm']), S=float(r['S_Nm']), TV=float(r['TV_Ta_Nm_per_s']), rev=float(r['Rev_pct']))
        elif r['Mean_eT_signed_pct_of_Tdref'] not in ('', 'NaN'):
            d = D.setdefault(c, {}); v = abs(float(r['Mean_eT_signed_pct_of_Tdref']))
            if not math.isnan(v): d['erel'] = max(d.get('erel', 0.0), v)
    return D


def collect():
    data = {}
    for c in CTRL:
        rows = [metrics(c, s) for s in SEEDS]
        missing = [s for s, r in zip(SEEDS, rows) if r is None]
        if missing:
            raise SystemExit('missing DREWS metrics for %s at seeds %s' % (c, missing))
        data[c] = rows
    return data


def ms(vals):
    return st.mean(vals), (st.stdev(vals) if len(vals) > 1 else 0.0)


def summary(data):
    S = {}
    for c in CTRL:
        for case in CASES:
            for k in ('rms', 'S', 'TV', 'rev', 'erel'):
                v = [r[case][k] for r in data[c] if k in r.get(case, {})]
                if v: S[(c, case, k)] = ms(v)
    return S


def paired(data, a, b, k):
    """mean and std over seeds of 100 (a - b) / b, per case, and the number of seeds where a < b"""
    out = {}
    for case in CASES:
        d = [100 * (ra[case][k] - rb[case][k]) / rb[case][k] for ra, rb in zip(data[a], data[b])]
        out[case] = (st.mean(d), st.stdev(d), sum(1 for x in d if x < 0))
    return out


# ------------------------------------------------------------------ figures
def fig_bars(S, k, ylabel, fname, log=False):
    fig, ax = plt.subplots(figsize=(9.5, 3.8))
    n = len(CTRL); w = 0.8 / n
    for i, c in enumerate(CTRL):
        xs = [j + (i - (n - 1) / 2) * w for j in range(len(CASES))]
        m = [S[(c, case, k)][0] for case in CASES]; s = [S[(c, case, k)][1] for case in CASES]
        ax.bar(xs, m, w, yerr=s, capsize=2, color=COLOR[c], edgecolor='k', linewidth=0.4, label=LABEL[c], hatch=None)
    ax.set_xticks(range(len(CASES))); ax.set_xticklabels(CASES); ax.set_ylabel(ylabel)
    if log: ax.set_yscale('log')
    ax.grid(axis='y', alpha=0.3); ax.legend(ncol=3, fontsize=8, frameon=False, loc='upper right')
    fig.tight_layout(); p = os.path.join(FIG, fname); fig.savefig(p, dpi=200); plt.close(fig); return p


def calib_pool(ctrl, pattern, seed, pmmin=30, gmmin=2):
    rows = []
    for f in glob.glob(os.path.join(RES, 'Controllers', ctrl, 'Calib', pattern)):
        for r in csv.DictReader(open(f, encoding='utf-8')):
            try:
                if int(float(r['seed'])) != seed or r['simulated'] != '1': continue
                if any(float(r[c]) > 0 for c in ('c1', 'c2', 'c3', 'c4')): continue
                if float(r['PM_deg']) < pmmin or float(r['GM']) < gmmin: continue
                rows.append((float(r['R_pct']), float(r['S_Nm'])))
            except (KeyError, ValueError):
                pass
    return rows


def front(pts):
    pts = sorted(pts, key=lambda p: p[1]); out = []; best = float('inf')
    for r, s in pts:
        if r < best: out.append((r, s)); best = r
    return out


def fig_pareto(designs):
    fig, ax = plt.subplots(figsize=(7.5, 4.6))
    pools = {'MapOpt': calib_pool('MapOpt', 'MapOpt_CALIB_*b_s00000.csv', 0), 'PIDOpt': calib_pool('PIDOpt', 'PIDOpt_CALIB_*_s10000.csv', 10000),
             'MapPIDOpt': calib_pool('MapPIDOpt', 'MapPIDOpt_CALIB_*_s30000.csv', 30000)}
    for c, pts in pools.items():
        if not pts: continue
        ax.scatter([p[1] for p in pts], [p[0] for p in pts], s=5, color=COLOR[c], alpha=0.25)
        f = front(pts); ax.plot([p[1] for p in f], [p[0] for p in f], color=COLOR[c], lw=1.4, label='%s: admissible evaluations and front' % LABEL[c])
    for c, (R, Sv) in designs.items():
        ax.scatter([Sv], [R], s=70, marker='o', color=COLOR[c], edgecolor='k', zorder=5, label='%s chosen' % LABEL[c])
    ax.set_xscale('log'); ax.set_yscale('log')
    ax.set_xlabel('S (RMS of $T_a$ above 5 Hz) [N·m]'); ax.set_ylabel('R (tracking error) [%]')
    ax.grid(alpha=0.3, which='both'); ax.legend(fontsize=7, frameon=False, ncol=1, loc='upper right')
    fig.tight_layout(); p = os.path.join(FIG, 'Opt_CALIB_pareto.png'); fig.savefig(p, dpi=200); plt.close(fig); return p


def sig(ctrl, case, seed=99999):
    p = os.path.join(RES, 'Controllers', ctrl, 'Cases', '%s_%s_sig_s%05d.csv' % (ctrl, case, seed))
    t, ts, td, e, ta = [], [], [], [], []
    for r in csv.DictReader(open(p, encoding='utf-8')):
        t.append(float(r['t'])); ts.append(float(r['T_s'])); td.append(float(r['T_d_ref'])); e.append(float(r['e_T'])); ta.append(float(r['T_a']))
    return t, ts, td, e, ta


def fig_time(case, fname):
    fig, axs = plt.subplots(3, 1, figsize=(9.5, 6.5), sharex=True)
    for c in ('MapOpt', 'PIDOpt', 'MapPIDOpt'):
        t, ts, td, e, ta = sig(c, case)
        axs[0].plot(t, ts, color=COLOR[c], lw=0.9, label=LABEL[c])
        axs[1].plot(t, e, color=COLOR[c], lw=0.9)
        axs[2].plot(t, ta, color=COLOR[c], lw=0.7)
    axs[0].plot(t, td, 'k--', lw=0.9, label='$T_{d,ref}$')
    axs[0].set_ylabel('$T_s$ [N·m]'); axs[1].set_ylabel('$e_T$ [N·m]'); axs[2].set_ylabel('$T_a$ [N·m]'); axs[2].set_xlabel('t [s]')
    for a in axs: a.grid(alpha=0.3)
    axs[0].legend(ncol=4, fontsize=8, frameon=False)
    fig.tight_layout(); p = os.path.join(FIG, fname); fig.savefig(p, dpi=200); plt.close(fig); return p


def fig_method():
    fig, ax = plt.subplots(figsize=(9.5, 3.2)); ax.axis('off')
    def box(x, y, w, h, txt, fc):
        ax.add_patch(matplotlib.patches.FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.02', fc=fc, ec='k', lw=0.6))
        ax.text(x + w / 2, y + h / 2, txt, ha='center', va='center', fontsize=7.5)
    def arrow(x1, y1, x2, y2):
        ax.annotate('', xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle='->', lw=0.8))
    ax.text(0.0, 0.93, 'Old (paper draft)', fontsize=9, weight='bold')
    old = ['Grid of parameters\n(Map: K x g_min, 61;\nPID: 824 points)', 'CALIB at ONE\nlocal seed', 'Gates: e_rel, stability\n(sim.), steering band', 'Min J = R x S\n(one point)', 'DREWS at ONE\nscoring seed 99999']
    for i, s in enumerate(old):
        box(0.0 + i * 0.2, 0.55, 0.17, 0.3, s, '#fde0c5')
        if i: arrow(i * 0.2 - 0.03, 0.70, i * 0.2, 0.70)
    ax.text(0.0, 0.43, 'New (common method, all controllers)', fontsize=9, weight='bold')
    new = ['Stage 1: search at seed x0000\n3 x (surrogateopt 250 +\npatternsearch 60), wide box', 'Stage 2: top 20\ndistinct admissible\n(PM>=30, GM>=2)', 'Stage 3: each on 5\nseeds x0001-x0005\nmin mean J (tie: PM)', 'Stage 4: calib_margin,\nbang-bang, rank\ncorrelation', 'Stage 5: DREWS on\n5 scoring seeds\nmean +- std, paired']
    for i, s in enumerate(new):
        box(0.0 + i * 0.2, 0.03, 0.17, 0.33, s, '#d0e8d0')
        if i: arrow(i * 0.2 - 0.03, 0.195, i * 0.2, 0.195)
    ax.set_xlim(-0.01, 0.98); ax.set_ylim(0, 1)
    fig.tight_layout(); p = os.path.join(FIG, 'Opt_method_flow.png'); fig.savefig(p, dpi=200); plt.close(fig); return p


# ------------------------------------------------------------------ docx helpers
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hexcolor); tcPr.append(s)


def table(doc, header, rows, widths=None, fs=8.5):
    t = doc.add_table(rows=1, cols=len(header)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(fs); shade(c, 'D9E2F3')
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(fs)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    return t


def caption(doc, text):
    p = doc.add_paragraph(); r = p.add_run(text); r.italic = True; r.font.size = Pt(9); p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead: p.add_run(bold_lead).bold = True
    p.add_run(text); p.paragraph_format.space_after = Pt(4); return p


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet'); p.add_run(text); p.paragraph_format.space_after = Pt(2)


def fig(doc, path, width_cm, cap):
    doc.add_picture(path, width=Cm(width_cm)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; caption(doc, cap)


def f(x, n=3):
    return ('%.' + str(n) + 'f') % x


# ------------------------------------------------------------------ main
def main():
    data = collect(); S = summary(data)
    dm, dp, dx = jload('map'), jload('pid'), jload('mappid')
    mo, po, xo = jload('mapopt'), jload('pidopt'), jload('mappidopt')
    # tables to csv
    with open(os.path.join(FIG, 'Opt_DREWS_summary_s99995-s99999.csv'), 'w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh); w.writerow(['controller', 'case', 'metric', 'mean', 'std'])
        for (c, case, k), (m, s) in sorted(S.items()): w.writerow([c, case, k, m, s])
    F1 = fig_method()
    F2 = fig_bars(S, 'rms', 'RMS $e_T$ [N·m] (log scale)', 'Opt_DREWS_rmse_5seeds.png', log=True)
    F3 = fig_bars(S, 'S', 'S [N·m]', 'Opt_DREWS_S_5seeds.png')
    F4 = fig_bars(S, 'TV', 'TV of $T_a$ [N·m/s]', 'Opt_DREWS_TV_5seeds.png')
    F5 = fig_pareto({'MapOpt': (mo['CALIB_R_pct_mean'], mo['CALIB_S_Nm_mean']), 'PIDOpt': (po['CALIB_Rbar_pct'], po['CALIB_Sbar_Nm']), 'MapPIDOpt': (xo['CALIB_R_pct_mean'], xo['CALIB_S_Nm_mean'])})
    F6 = fig_time('SLIP', 'Opt_SLIP_time_s99999.png')

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(2.0)
    st_ = doc.styles['Normal']; st_.font.name = 'Calibri'; st_.font.size = Pt(10.5)
    doc.add_heading('Common parameter-selection method for the C-EPS controllers: method and first comparison', 0)
    para(doc, 'Trần Minh Tân — progress report to TS. Phùng Thanh Huy — 2026-10-17')

    doc.add_heading('1. Summary', 1)
    bullet(doc, 'All controllers are now tuned with ONE common procedure: a global search (surrogateopt, then patternsearch) on the calibration case CALIB at one seed, '
                'the 20 best distinct admissible designs re-scored on five other seeds, and the design with the smallest MEAN J = R × S chosen. The old grid sweeps '
                'with one seed are kept as references.')
    bullet(doc, 'Results are given for the three controllers tuned with this method: Map (assist map with lead-lag), PID and MapPID (map branch + PID branch), '
                'scored on the five test cases DRY, RAMP, SLIP, WEAVE, EDGE at five scoring seeds (99995-99999), as mean ± standard deviation over the seeds.')
    rp = paired(data, 'MapPIDOpt', 'PIDOpt', 'rms'); sp = paired(data, 'MapPIDOpt', 'PIDOpt', 'S')
    rm = paired(data, 'MapPIDOpt', 'MapOpt', 'rms'); sm = paired(data, 'MapPIDOpt', 'MapOpt', 'S')
    bullet(doc, 'MapPID vs PID (same tuning effort): RMS e_T %s %%, S %s %% (DRY, RAMP, SLIP, WEAVE, EDGE).'
           % (', '.join('%+.0f' % rp[c][0] for c in CASES), ', '.join('%+.0f' % sp[c][0] for c in CASES)))
    bullet(doc, 'MapPID vs Map: RMS e_T %s %%, S %s %%.' % (', '.join('%+.0f' % rm[c][0] for c in CASES), ', '.join('%+.0f' % sm[c][0] for c in CASES)))
    doc.add_heading('2. Old method (used in the paper draft)', 1)
    para(doc, 'Each controller was tuned on CALIB (117 s: steady cornering 0.1–0.3 g at 20/60/100 km/h, sine steering, small corrections and a friction step) with '
              'ONE noise seed of its own range. R is the tracking error (mean of four group errors in % of T_d,ref), S the RMS of the assist torque above 5 Hz '
              '(noise), J = R × S. Gates: steady-state error e_rel ≤ 5 % on 9 windows, stability with margin by simulation (gain × 2 and +1 ms delay), '
              'steering-band constraint (filter gain ≥ 0.707 at 2 Hz). The admissible point with the smallest J was chosen.')
    bullet(doc, 'Map: lead-lag designed off line on 68 linear plants (smallest high-frequency gain with PM ≥ 45°), only the slope limit K simulated (61 configurations).')
    bullet(doc, 'PID: grid on Kp, Ki, Kd, Tf (824 configurations).')
    bullet(doc, 'MapPID: first simultaneous optimisation of 9 parameters (930 evaluations) at one seed, checked on one more seed.')
    para(doc, 'Weak points: different tuning effort and method per controller (so a better result can come from the method, not the structure), narrow boxes '
              '(MapPID stuck at the lower bound of its lead), and one-seed selection (the ranking of points closer than about 3 % is noise).')

    doc.add_heading('3. New common method', 1)
    fig(doc, F1, 16.5, 'Figure 1. Old and new parameter-selection procedures.')
    rows = [['Parameters searched', 'grid (Map: K only, lead-lag off line)', 'all parameters at once, surrogateopt ×3 + patternsearch ×3 (≈ 930 evaluations)'],
            ['Box', 'narrow, hand set', 'wide, log scale; lead p/z ≥ 1 and lag p1/z1 ≤ 1 (stages can switch off); K 0.1–16; Kp 1e-3–30; Ki 0.1–2000; Kd 0 or 1e-4–2; Tf 1–79.6 ms'],
            ['Linear requirement', 'Map PM ≥ 45° (design), others report', 'PM ≥ 30°, GM ≥ 2 on 68 plants (also during the search for Map and MapPID)'],
            ['Seeds for the choice', '1', 'search at x0000; choice on 5 seeds x0001–x0005 (x0000 not in the mean)'],
            ['Choice', 'min J at one seed', 'min mean J over 5 seeds among the 20 best; tie (difference < std) → larger PM; e_rel ≤ 5 % at every seed'],
            ['Checks', 'calib_margin', 'calib_margin, bang-bang, rank correlation search seed vs mean, convergence of 3 chains'],
            ['Test cases', '1 scoring seed', '5 scoring seeds, mean ± std, paired differences']]
    table(doc, ['Item', 'Old', 'New'], rows, widths=[3.2, 5.0, 8.8])
    caption(doc, 'Table 1. Differences between the two methods.')
    para(doc, 'Lesson learned while applying it: the linear requirement must already be imposed during the search. For the assist map the best J lies at PM ≈ 15°; '
              'a search with PM > 0 and a filter PM ≥ 30° afterwards left only 7 admissible designs of 583. Re-run with PM ≥ 30° in the search, Map reached '
              'J = %.2f instead of 6.29.' % mo['CALIB_J_mean_seeds_1_5'])

    doc.add_heading('4. Chosen designs', 1)
    rows = [['Map', 'K %.2f, lead z %.1f p %.1f, lag z1 %.1f p1 %.1f' % (mo['Kmax'], mo['lead']['z'], mo['lead']['p'], mo['lag']['z1'], mo['lag']['p1']), '%s ± %s' % (f(mo['CALIB_J_mean_seeds_1_5']), f(mo['CALIB_J_sigma'])), '%.2f / %.3f' % (mo['CALIB_R_pct_mean'], mo['CALIB_S_Nm_mean']), f(mo['PM_deg'], 1), f(mo['GM'], 2)],
            ['PID', 'Kp %.2f, Ki %.1f, Kd %.3f, Tf %.1f ms' % (po['Kp'], po['Ki'], po['Kd'], 1e3 * po['Tf']), '%s ± %s' % (f(po['CALIB_Jbar']), f(po['CALIB_sigmaJ'])), '%.2f / %.3f' % (po['CALIB_Rbar_pct'], po['CALIB_Sbar_Nm']), f(po['PM_deg'], 1), f(po['GM'], 2)],
            ['MapPID', 'K %.2f, lead z %.0f p %.0f, lag z1 %.1f p1 %.1f, Kp %.3f, Ki %.1f, Kd %.3f, Tf %.1f ms' % (xo['Kmax'], xo['lead']['z'], xo['lead']['p'], xo['lag']['z1'], xo['lag']['p1'], xo['Kp'], xo['Ki'], xo['Kd'], 1e3 * xo['Tf']), '%s ± %s' % (f(xo['CALIB_J_mean']), f(xo['CALIB_J_sigma'])), '%.2f / %.3f' % (xo['CALIB_R_pct_mean'], xo['CALIB_S_Nm_mean']), f(xo['PM_deg'], 1), f(xo['GM'], 2)]]
    table(doc, ['Controller', 'Parameters (rad/s, N·m, s)', 'CALIB J (5 seeds)', 'R [%] / S [N·m]', 'PM [deg]', 'GM'], rows, widths=[2.0, 7.6, 2.8, 2.6, 1.3, 1.1])
    caption(doc, 'Table 2. Chosen designs: mean ± std over the 5 choice seeds; PM and GM over 68 linear plants.')
    para(doc, 'Observations: in MapPID the lead is practically switched off (p/z = %.2f, corner at %.0f rad/s), the lag stays, and Kp ≈ 0 (the map branch supplies the '
              'proportional action). Within the top 20 of MapPID the J differences (0.03–0.05) are smaller than the seed-to-seed spread; the rank correlation '
              'between the search seed and the 5-seed mean is only %.2f (PID: %.2f), which is why the choice must use several seeds.' % (xo['lead']['p'] / xo['lead']['z'], xo['lead']['z'], xo['rank_correlation'], po['rank_correlation']))
    fig(doc, F5, 14.5, 'Figure 2. CALIB tracking–noise plane: admissible evaluations of the -Opt searches with their fronts, and the chosen designs '
                      '(circles: mean over the 5 choice seeds). Map, PID and MapPID were searched at different seeds (0, 10000, 30000).')

    doc.add_heading('5. Results on the test cases (5 scoring seeds)', 1)
    for k, name, n in (('rms', 'RMS tracking error e_T [N·m]', 3), ('S', 'noise S [N·m]', 3), ('TV', 'total variation of T_a [N·m/s]', 1), ('erel', 'largest steady-state error |e_rel| [%]', 2), ('rev', 'reverse assist [% of time]', 1)):
        cases = [c for c in (('DRY', 'RAMP', 'EDGE') if k == 'erel' else CASES) if all((ct, c, k) in S for ct in CTRL)]   # e_rel only on the steady-cornering cases
        rows = [[LABEL[ct]] + ['%s ± %s' % (f(S[(ct, c, k)][0], n), f(S[(ct, c, k)][1], n)) for c in cases] for ct in CTRL]
        table(doc, ['Controller'] + cases, rows, widths=[3.0] + [2.8] * len(cases), fs=8)
        caption(doc, 'Table: %s, mean ± std over seeds 99995–99999%s.' % (name, ' (steady windows of DRY, RAMP, EDGE)' if k == 'erel' else ''))
    fig(doc, F2, 16.5, 'Figure 3. RMS e_T per case (mean ± std over 5 seeds, log scale).')
    fig(doc, F3, 16.5, 'Figure 4. Noise S per case (mean ± std over 5 seeds).')
    fig(doc, F4, 16.5, 'Figure 5. Total variation of the assist torque per case.')
    fig(doc, F6, 16.5, 'Figure 6. SLIP case (friction changes while cornering), seed 99999: T_s with T_d,ref, error e_T and assist T_a of the three controllers.')

    doc.add_heading('6. Paired comparisons', 1)
    para(doc, 'The assist map is almost insensitive to the sensor noise: its standard deviation over seeds is below 0.001 N·m and shows as 0.000.')
    para(doc, 'Relative difference 100 (A − B)/B per seed, mean ± std over the 5 seeds, and the number of seeds where A is lower (better).')
    pairs = [('MapPIDOpt', 'PIDOpt'), ('MapPIDOpt', 'MapOpt'), ('PIDOpt', 'MapOpt')]
    for k, name in (('rms', 'RMS e_T'), ('S', 'S')):
        rows = []
        for a, b in pairs:
            d = paired(data, a, b, k)
            rows.append(['%s vs %s' % (LABEL[a], LABEL[b])] + ['%+.0f ± %.0f (%d/5)' % d[c] for c in CASES])
        table(doc, ['Pair (A vs B)'] + CASES, rows, widths=[4.2] + [2.56] * 5, fs=8)
        caption(doc, 'Table: paired difference of %s [%%].' % name)

    doc.add_heading('7. Discussion', 1)
    bullet(doc, 'Same method, different structure: MapPID vs PID isolates the effect of the map branch on top of a PID tuned with the same effort.')
    bullet(doc, 'In MapPID the optimiser removed the lead (p/z close to 1) and set Kp close to 0: the map branch takes over the proportional action and the PID '
                'branch keeps the integral and derivative parts; the lag on the map branch remains. CALIB cannot distinguish designs with different Kp '
                '(their group errors and S are equal), while the speed ramp, friction changes and weaving of RAMP, SLIP and WEAVE do; this is a limitation of '
                'J = R × S on CALIB.')
    bullet(doc, 'The assist map alone keeps the lowest noise and almost no reverse assist but large steady-state errors (no integral action), whatever the tuning.')
    doc.add_heading('8. Limitations and next steps', 1)
    bullet(doc, 'One search seed per controller and five choice seeds; the -Opt searches are local optima (convergence checked by 3 chains).')
    bullet(doc, 'The rule was changed after seeing results (common method, PM ≥ 30° during the search); this must be stated in the paper.')
    bullet(doc, 'Possible next step (needs a decision): add a dynamic test (speed ramp) or the total variation to the selection criterion, so that CALIB sees the effect of Kp.')
    doc.save(OUT)
    print('written', OUT)


if __name__ == '__main__':
    main()
