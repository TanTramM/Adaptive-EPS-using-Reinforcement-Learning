"""Refresh Documents/Map/Map.xlsx and Documents/Ref/Reference.xlsx from the model data and results.

Usage:  python tools/export_xlsx.py        (close both Excel files first; run after make_ta_max / calibrate_map / export_map_results)

Sources:
  Model/data/map.json                          map table, parameters, lead compensator, calibration pairs
  Model/data/ref.json  (field Ta_max)          assist limit T_a,max(v) of Documents/Ref/ref.txt section 1.5
  Result/Map/Calibration, OverAssist, TestCases  dry-road error, over-assist, test-case metrics
Only the values of the existing sheets are replaced; titles, widths and styles of the header row are kept.
"""
import csv
import json
import os

import numpy as np
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def path(*p):
    return os.path.join(ROOT, *p)


def rd_csv(*p):
    return list(csv.reader(open(path(*p), encoding='utf-8')))


def clear(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)


m = json.load(open(path('Model', 'data', 'map.json'), encoding='utf-8'))
ref = json.load(open(path('Model', 'data', 'ref.json'), encoding='utf-8'))
ta = ref['Ta_max']
v = m['v_breakpoints_kmh']
assert v == ta['v_kmh'], 'map.json and ref.json Ta_max speeds differ'

# ---------------- Map.xlsx ----------------
p = path('Documents', 'Map', 'Map.xlsx')
wb = openpyxl.load_workbook(p)

ws = wb['Bản đồ T_a']
clear(ws)
tab = np.array(m['Ta_table_Nm'])
for k, t in enumerate(m['Ts_breakpoints_Nm']):
    ws.append([t] + [float(x) for x in tab[:, k]])

ws = wb['Tham số']
clear(ws)
L, H = m['lead'], m['lead_hi']
fmt = lambda x: '[' + ' '.join('%.6f' % q for q in x) + ']'
for row in [
        ['T_ctl', m['Ts_ctrl']['value'], 's', 'Chu kỳ lấy mẫu của vòng trợ lực', 'Tự chọn(8)'],
        ['T_s0', m['Ts0']['value'], 'N.m', 'Vùng chết của bản đồ', 'Tự chọn(6)'],
        ['K_max bậc thấp (20-60 km/h)', L['K_max'], '-', 'Độ dốc lớn nhất của bản đồ ở vận tốc thấp', 'Tự chọn(7)'],
        ['K_max bậc cao (65-100 km/h)', H['K_max'], '-', 'Độ dốc lớn nhất của bản đồ ở vận tốc cao', 'Tự chọn(7)'],
        ['Vùng trộn hai khâu bù [km/h]', '%g - %g' % tuple(m['blend_kmh']), 'km/h', 'Trọng số của khâu bậc cao tăng tuyến tính từ 0 lên 1', 'Tự chọn(7)'],
        ['z_l bậc thấp', L['zero_rad_s'], 'rad/s', 'Điểm không của khâu lead bậc thấp', 'Thiết kế theo tiêu chí của [1]'],
        ['p_l bậc thấp', L['pole_rad_s'], 'rad/s', 'Điểm cực của khâu lead bậc thấp', 'Thiết kế theo tiêu chí của [1]'],
        ['z_l bậc cao', H['zero_rad_s'], 'rad/s', 'Điểm không của khâu lead bậc cao', 'Thiết kế theo tiêu chí của [1]'],
        ['p_l bậc cao', H['pole_rad_s'], 'rad/s', 'Điểm cực của khâu lead bậc cao', 'Thiết kế theo tiêu chí của [1]'],
        ['Độ dự trữ pha nhỏ nhất bậc thấp', round(L['min_phase_margin_deg'], 1), 'độ', 'Với K_v từ 0.5 tới K_max, k_r đường khô và k_r = 0, có motor', 'Tính toán'],
        ['Độ dự trữ pha nhỏ nhất bậc cao', round(H['min_phase_margin_deg'], 1), 'độ', 'Như trên', 'Tính toán'],
        ['Độ dự trữ biên nhỏ nhất bậc thấp', round(L['min_gain_margin'], 2), 'lần', 'Như trên', 'Tính toán'],
        ['Độ dự trữ biên nhỏ nhất bậc cao', round(H['min_gain_margin'], 2), 'lần', 'Như trên', 'Tính toán'],
        ['Hệ số tử số H(z) bậc thấp', fmt(L['num']), '-', 'Một khâu lead rời rạc (Tustin)', 'Tính toán'],
        ['Hệ số mẫu số H(z) bậc thấp', fmt(L['den']), '-', 'Một khâu lead rời rạc (Tustin)', 'Tính toán'],
        ['Hệ số tử số H(z) bậc cao', fmt(H['num']), '-', 'Một khâu lead rời rạc (Tustin)', 'Tính toán'],
        ['Hệ số mẫu số H(z) bậc cao', fmt(H['den']), '-', 'Một khâu lead rời rạc (Tustin)', 'Tính toán'],
        ['ω_m (motor)', round(L['motor_wm_rad_s'], 1), 'rad/s', 'Tần số cắt của khâu trễ motor (100 Hz)', '[1] (Bảng III)'],
        ['μ hiệu chỉnh', m['calibration_mu'], '-', 'Hệ số bám khi hiệu chỉnh bản đồ', '[4]']]:
    ws.append(row)

ws = wb['T_a,max']
clear(ws)
for i in range(len(v)):
    ws.append([v[i], round(ta['a_y_max_g'][i], 4), round(ta['a_y_lim_g'][i], 4), round(ta['T_r_top_Nm'][i], 4),
               round(ta['T_d_ref_top_Nm'][i], 4), round(ta['Ta_max_Nm'][i], 4)])

ws = wb['Cặp hiệu chỉnh']
clear(ws)
for i, c in enumerate(m['calibration']):
    for a, t, tA, tm in zip(c['ay_g'], c['Ts_Nm'], c['Ta_Nm'], c['Ta_map_Nm']):
        ws.append([v[i], a, round(t, 4), round(tA, 4), round(tm, 4)])

for name, rel in [('Sai số đường khô', ('Result', 'Map', 'Calibration', 'Map_dry_road_steady_error.csv')),
                  ('Trợ lực thừa', ('Result', 'Map', 'OverAssist', 'Map_over_assist_vs_lateral_accel.csv'))]:
    ws = wb[name]
    clear(ws)
    for row in rd_csv(*rel)[1:]:
        ws.append([float(x) for x in row])

ws = wb['Ca thử TC1-TC6']
clear(ws)
for r in rd_csv('Result', 'Map', 'TestCases', 'Map_test_case_metrics.csv')[1:]:
    ws.append([r[0], r[1]] + [None if x == 'NaN' else float(x) for x in r[2:]])
wb.save(p)
print('Map.xlsx saved')

# ---------------- Reference.xlsx (tab T_a,max) ----------------
p2 = path('Documents', 'Ref', 'Reference.xlsx')
wb2 = openpyxl.load_workbook(p2)
ws = wb2['T_a,max']
for i in range(len(v)):
    r = i + 2
    assert ws.cell(r, 1).value == v[i]
    ws.cell(r, 6).value = round(ta['T_r_top_Nm'][i], 3)
    ws.cell(r, 7).value = round(ta['T_d_ref_top_Nm'][i], 3)
    ws.cell(r, 8).value = round(ta['Ta_max_Nm'][i], 3)
wb2.save(p2)
print('Reference.xlsx saved')
