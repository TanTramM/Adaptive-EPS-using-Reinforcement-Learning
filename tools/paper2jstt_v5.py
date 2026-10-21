"""Create the JSTT submission manuscript in the v5/v6 one-column layout.

Usage:
  python tools/paper2jstt_v5.py Documents/Paper/paper_merged_v7.txt Documents/Paper/paper_v7_sections1-3.docx --through 3
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, end])


def configure(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(0)
    for name, size in [('Title', 15), ('Heading 1', 13), ('Heading 2', 12)]:
        style = doc.styles[name]; style.font.name = 'Times New Roman'; style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        style.font.size = Pt(size); style.font.bold = True
        style.paragraph_format.line_spacing = 1.5; style.paragraph_format.space_before = Pt(10); style.paragraph_format.space_after = Pt(4)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_page_number(footer)


def read_source(path):
    lines = open(path, encoding='utf-8').read().splitlines()
    bars = [i for i, x in enumerate(lines) if re.fullmatch(r'={20,}', x.strip())]
    if len(bars) < 2:
        raise ValueError('Source must begin with the title between two = separators.')
    title = lines[bars[0] + 1].strip()
    start = next(i for i, x in enumerate(lines) if x.strip() == 'Abstract')
    keywords_i = next(i for i, x in enumerate(lines[start + 1:], start + 1) if x.startswith('Keywords:'))
    abstract = ' '.join(x.strip() for x in lines[start + 1:keywords_i] if x.strip() and not re.fullmatch(r'-{20,}', x.strip()))
    return title, abstract, lines[keywords_i][len('Keywords:'):].strip(), lines[keywords_i + 1:]


def make_para(doc, text='', style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True; p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_table(doc, caption, rows):
    cap = make_para(doc, caption); cap.runs[0].bold = True
    table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    table.style = 'Table Grid'
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            p = cell if False else table.cell(ri, ci).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(cell); r.font.name = 'Times New Roman'; r.font.size = Pt(10); r.bold = ri == 0


def add_figure(doc, caption, spec):
    file_name, _, width = spec.partition('@')
    full = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), file_name)
    if os.path.isfile(full):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(full, width=Cm(float(width or 14)))
    else:
        p = make_para(doc, '[Insert figure here: ' + file_name + ']'); p.runs[0].italic = True
    cap = make_para(doc, caption); cap.runs[0].bold = True


def build(src, dst, through=None):
    title, abstract, keywords, lines = read_source(src)
    doc = Document(); configure(doc)
    p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(title)
    make_para(doc, 'Abstract: ' + abstract, bold_prefix='Abstract: ')
    make_para(doc, 'Keywords: ' + keywords, bold_prefix='Keywords: ')
    i = 0
    while i < len(lines):
        raw, s = lines[i], lines[i].strip()
        if not s or re.fullmatch(r'[-=]{20,}', s): i += 1; continue
        h1 = re.fullmatch(r'(\d+)\.\s+(.+)', s)
        h2 = re.fullmatch(r'(\d+)\.(\d+)\.\s+(.+)', s)
        if h1 and not h2:
            if through and int(h1.group(1)) > through: break
            make_para(doc, s, style='Heading 1'); i += 1; continue
        if h2:
            make_para(doc, s, style='Heading 2'); i += 1; continue
        eq = re.fullmatch(r'\((\d+)\)\s+(.+)', s)
        if eq and i + 1 < len(lines) and lines[i + 1].strip().startswith('['):
            p = make_para(doc, eq.group(2) + '    (' + eq.group(1) + ')'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 2; continue
        fig = re.fullmatch(r'\[(Fig\.\s+\d+\.\s+.+?)\s+\|\s+(.+)\]', s)
        if fig:
            add_figure(doc, fig.group(1), fig.group(2)); i += 1; continue
        if s.startswith('Table ') and i + 1 < len(lines) and ' | ' in lines[i + 1]:
            rows = []; j = i + 1
            while j < len(lines) and ' | ' in lines[j]: rows.append([x.strip() for x in lines[j].split(' | ')]); j += 1
            add_table(doc, s, rows); i = j; continue
        if s.startswith('- '):
            make_para(doc, '• ' + s[2:]); i += 1; continue
        make_para(doc, s); i += 1
    doc.save(dst)
    print('Wrote', dst)


if __name__ == '__main__':
    if len(sys.argv) < 3: raise SystemExit('Usage: paper2jstt_v5.py input.txt output.docx [--through N]')
    limit = int(sys.argv[sys.argv.index('--through') + 1]) if '--through' in sys.argv else None
    # Use the advisor-derived generator: it converts $...$ and display LaTeX to OMML.
    import paper2docx
    paper2docx.build(sys.argv[1], sys.argv[2], limit)
