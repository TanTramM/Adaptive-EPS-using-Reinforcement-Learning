"""Build the journal manuscript (JSTT, University of Transport Technology) from Documents/Paper/paper_jstt*.txt on the journal's own Word template.

Usage:  python tools/paper2jstt.py Documents/Paper/paper_jstt_v1.txt Documents/Paper/paper_jstt_v1.docx

The template is References/JSTT_Manuscript_template_EN.docx: first-page banner and the title block (a floating table: title, authors, affiliations, Article
info, Abstract, Keywords), then a two-column body (Arial 11 pt, styles '10.JSTT-Text', '4.JSST-1storder-head', '4.1JSTT-2ndorder-head',
'JSTT-acknowledgement', 'JSTT-table-caption', 'JSTT-figure-caption', 'EndNote Bibliography'). Headings are numbered by their styles, so the numbers of the .txt
are removed. Wide objects (figures, tables, long equations) sit in one-column continuous sections between two-column text. Equations are Word equations
numbered at the right margin; tables have no vertical rules; figures come from Figures/v5_Huypt (PNG 300 dpi; the vector PDF and 1000 dpi TIFF copies
for the submission are in Figures/JSTT).

Author data are filled from the AUTHORS block below; what the authors still have to supply is written in square brackets.
"""
import os
import re
import sys
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import txt2docx
txt2docx.AUTO_PAREN = True
from paper2docx import add_rich, equation
from txt2docx import REPO_ROOT, latex_to_omath

TEMPLATE = os.path.join(REPO_ROOT, 'References', 'JSTT_Manuscript_template_EN.docx')
COL_W = 8.4                         # [cm] width of one column
TEXT_W = 17.4                       # [cm] text width of the template (page 21 cm, margins 1.83 and 1.70 cm)
MAX_FIG_H = 12.0                    # [cm]
FLOW = bool(os.environ.get('JSTT_FLOW'))   # no pictures, no one-column blocks: tables stay inside the column, text runs on column after column
WIDE_EQ = set()                     # every display equation fits one column (long ones are broken into rows in the .txt)
UNNUMBERED_H = ('Acknowledgments', 'Author contributions', 'Conflict of Interest', 'Declaration of Generative AI', 'References')

AUTHORS = [('Tran Minh Tan', '1,*'), ('Phung Thanh Huy', '1')]
AFFILIATION = ('1', 'Ho Chi Minh City University of Technology (HCMUT), Vietnam National University Ho Chi Minh City (VNU-HCM), '
                    '[address, city, postcode], Vietnam')
CORRESPONDING = 'Tran Minh Tan'
EMAIL = '[e-mail to be added]'
ABBREVIATIONS = ('Abbreviations: C-EPS, column-type electric power steering; EPS, electric power steering; PI, proportional–integral; '
                 'PID, proportional–integral–derivative; RMS, root mean square; CALIB, calibration scenario; DRY, dry-road scenario; '
                 'SLIP, friction-change scenario; WEAVE, continuous-steering scenario; RAMP, speed-ramp scenario; EDGE, high-lateral-acceleration scenario. '
                 'Units: lateral acceleration in g (1 g = 9.81 m/s²); speed in km/h (1 km/h = 0.278 m/s).')


# ------------------------------------------------------------------------------------------------ small helpers
def set_par(p, parts):
    """replace the content of paragraph p by runs [(text, {'b':..,'i':..,'sup':..}), ...] keeping its style"""
    for ch in list(p._p):
        if ch.tag != qn('w:pPr'):
            p._p.remove(ch)
    for text, fmt in parts:
        r = p.add_run(text)
        if fmt.get('b'):
            r.bold = True
        if fmt.get('i'):
            r.italic = True
        if fmt.get('sup'):
            r.font.superscript = True


def remove_par(p):
    p._p.getparent().remove(p._p)


def set_size(p, pt):
    """font size of the runs and of the inline equations of paragraph p"""
    for r in p.runs:
        r.font.size = Pt(pt)
    for mr in p._p.iter(qn('m:r')):
        rpr = mr.find(qn('w:rPr'))
        if rpr is None:
            rpr = OxmlElement('w:rPr')
            mrpr = mr.find(qn('m:rPr'))
            mr.insert(0 if mrpr is None else 1, rpr)
        for tag in ('w:sz', 'w:szCs'):
            if rpr.find(qn(tag)) is None:
                el = OxmlElement(tag); el.set(qn('w:val'), str(int(2 * pt))); rpr.append(el)


def set_cols(sectPr, n):
    cols = sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols'); sectPr.append(cols)
    cols.set(qn('w:num'), str(n)); cols.set(qn('w:space'), '360')


def strip_title_pg(sectPr):
    for el in sectPr.findall(qn('w:titlePg')):
        sectPr.remove(el)


def shrink_break_paragraph(p_el):
    """the paragraph that carries a section break: make it almost invisible"""
    pPr = p_el.find(qn('w:pPr'))
    sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '0'); sp.set(qn('w:line'), '20'); sp.set(qn('w:lineRule'), 'exact')
    pPr.insert(0, sp)
    rpr = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '2'); rpr.append(sz)
    pPr.append(rpr)


class Layout:
    """two-column text with one-column blocks; section breaks are created on demand"""
    def __init__(self, doc):
        self.doc = doc; self.cols = 2; self.breaks = 0

    def to(self, n):
        if n == self.cols:
            return
        sec = self.doc.add_section(WD_SECTION.CONTINUOUS)
        body = self.doc.element.body
        brk = body[-2]                                    # the paragraph that now ends the previous section
        shrink_break_paragraph(brk)
        self.breaks += 1
        if self.breaks > 1:
            strip_title_pg(brk.find(qn('w:pPr')).find(qn('w:sectPr')))
        set_cols(sec._sectPr, n); strip_title_pg(sec._sectPr)
        self.cols = n


# ------------------------------------------------------------------------------------------------ front page
def fill_front(doc, title, abstract, keywords):
    tbl = doc.tables[0]
    c = tbl.cell(0, 2)
    ps = c.paragraphs
    set_par(ps[0], [(title, {})])
    parts = []
    for i, (name, aff) in enumerate(AUTHORS):
        parts.append((name, {})); parts.append((aff, {'sup': True}))
        if i < len(AUTHORS) - 1:
            parts.append((', ', {}))
    set_par(ps[1], parts)
    set_par(ps[2], [(AFFILIATION[0], {'sup': True}), (AFFILIATION[1], {})])
    for extra in ps[3:]:
        remove_par(extra)
    info = tbl.cell(1, 0).paragraphs
    set_par(info[4], [('Corresponding author: ' + CORRESPONDING, {})])
    set_par(info[6], [(EMAIL, {})])
    set_par(info[8], [('Received: dd Month yyyy', {})])
    set_par(info[9], [('Accepted: dd Month yyyy', {})])
    set_par(info[10], [('Published: dd Month yyyy', {})])
    ab = tbl.cell(1, 2).paragraphs
    set_par(ab[0], [('Abstract: ', {'b': True}), (abstract, {})])
    set_par(ab[-1], [('Keywords: ', {'b': True}), (keywords, {})])
    for p in ab[1:-1]:
        remove_par(p)


# ------------------------------------------------------------------------------------------------ blocks
def add_text(doc, style, text, indent0=False, size=None, bold=False, italic=False, align=None):
    p = doc.add_paragraph(style=style)
    add_rich(p, text, bold=bold, italic=italic)
    if indent0:
        p.paragraph_format.first_line_indent = Cm(0)
    if size:
        set_size(p, size)
    if align is not None:
        p.alignment = align
    return p


def jstt_equation(doc, latex, num):
    """display equation numbered at the right of the column; rows separated by a double backslash are the lines of a broken long equation"""
    rows = [r.strip() for r in latex.split('\\\\')]
    p = doc.add_paragraph(style='10.JSTT-Text'); p.paragraph_format.first_line_indent = Cm(0)
    para = txt2docx.m('oMathPara'); om = txt2docx.msub(para, 'oMath'); arr = txt2docx.msub(om, 'eqArr')
    pr = txt2docx.msub(arr, 'eqArrPr'); md = txt2docx.msub(pr, 'maxDist'); md.set('{%s}val' % txt2docx.M_NS, '1')
    for i, r in enumerate(rows):
        e = txt2docx.msub(arr, 'e')
        e.extend(list(latex_to_omath(r)))
        if i == len(rows) - 1:
            e.append(txt2docx.run('#', upright=True))
            d = txt2docx.msub(e, 'd'); de = txt2docx.msub(d, 'e'); de.append(txt2docx.run(str(num), upright=True))
    for r in para.iter('{%s}r' % txt2docx.M_NS):
        rpr = OxmlElement('w:rPr')
        for tag in ('w:sz', 'w:szCs'):
            el = OxmlElement(tag); el.set(qn('w:val'), '20'); rpr.append(el)
        mrpr = r.find('{%s}rPr' % txt2docx.M_NS)
        r.insert(0 if mrpr is None else 1, rpr)
    p._p.append(para)


NS_DECL = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
           'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
           'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
           'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
           'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"')
CAP_PT = 9.0
_group_id = [100]


def figure_group(doc, items, caption, total_w_cm):
    """w:drawing holding ONE inline group: the pictures (one under the other, each with its panel label in a small borderless text box) and, at the bottom, a
    borderless text box with the caption; pictures and caption move together.  items = [(file, width_cm, height_cm, panel_label or '')]"""
    from docx.text.paragraph import Paragraph
    from docx.oxml import parse_xml
    emu = lambda cm: int(round(cm * 360000))
    gap = emu(0.12)
    cw = emu(total_w_cm)
    _group_id[0] += 1
    n0 = _group_id[0] * 10
    y = 0
    shapes = []
    boxes = []                                                # (xml placeholder id, text, size pt, align, bold)
    k = 0
    for (full, wcm, hcm, label) in items:
        rid, _ = doc.part.get_or_add_image(full)
        pw, ph = emu(wcm), emu(hcm)
        x = (cw - pw) // 2
        k += 1
        shapes.append(
            f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{n0 + k}" name="Picture {n0 + k}"/><pic:cNvPicPr/></pic:nvPicPr>'
            f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
            f'<pic:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{pw}" cy="{ph}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>')
        y += ph + gap
        if label:
            lh = emu(0.50)
            k += 1
            shapes.append(_textbox_xml(n0 + k, 0, y, cw, lh))
            boxes.append((label, CAP_PT + 0.5, True))
            y += lh + gap
    plain = re.sub(r'\$[^$]*\$', 'xx', caption)
    per_line = max(10, int(total_w_cm / 0.158))               # characters of 9 pt Arial per line
    lines = -(-len(plain) // per_line)
    cap_h = emu(0.20 + 0.40 * lines)
    k += 1
    shapes.append(_textbox_xml(n0 + k, 0, y, cw, cap_h))
    boxes.append((caption, CAP_PT, False))
    total_h = y + cap_h
    xml = (
        f'<w:drawing {NS_DECL}><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="{cw}" cy="{total_h}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:docPr id="{n0}" name="Figure group {n0}"/><wp:cNvGraphicFramePr/>'
        f'<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"><wpg:wgp><wpg:cNvGrpSpPr/>'
        f'<wpg:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cw}" cy="{total_h}"/><a:chOff x="0" y="0"/><a:chExt cx="{cw}" cy="{total_h}"/></a:xfrm></wpg:grpSpPr>'
        + ''.join(shapes) + '</wpg:wgp></a:graphicData></a:graphic></wp:inline></w:drawing>')
    el = parse_xml(xml)
    for content, (text, size, bold) in zip(el.iter(qn('w:txbxContent')), boxes):
        p_el = OxmlElement('w:p')
        content.append(p_el)
        par = Paragraph(p_el, doc._body)
        par.style = doc.styles['JSTT-figure-caption']
        add_rich(par, text, bold=bold)
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.space_before = Pt(0); par.paragraph_format.space_after = Pt(0)     # no style spacing: the text box has its own height
        par.paragraph_format.line_spacing = 1.0
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_size(par, size)
    return el


def _textbox_xml(nid, x, y, w, h):
    return (f'<wps:wsp><wps:cNvPr id="{nid}" name="Text box {nid}"/><wps:cNvSpPr txBox="1"/>'
            f'<wps:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
            f'<a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr><wps:txbx><w:txbxContent/></wps:txbx>'
            f'<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" anchorCtr="0"><a:noAutofit/></wps:bodyPr></wps:wsp>')


def add_figure(doc, caption, path, wide=True):
    """path = one file, or several pictures separated by ' ; ', each 'file@width_cm|panel label' (inserted as separate pictures in ONE group with the caption)"""
    items = []
    for part in [p.strip() for p in path.split(' ; ')]:
        spec, _, label = part.partition('|')
        file_, _, wcm = spec.partition('@')
        full = os.path.join(REPO_ROOT, file_.strip())
        w, h = Image.open(full).size
        width = float(wcm) if wcm else (min(TEXT_W, MAX_FIG_H * w / h) if wide else min(COL_W, 20.0 * w / h))
        width = min(width, TEXT_W)
        items.append((full, width, width * h / w, label.strip()))
    total_w = max(i[1] for i in items) if len(items) == 1 else TEXT_W if wide else COL_W
    box = doc.add_paragraph(style='0.đánh dấu hình'); box.paragraph_format.keep_with_next = True
    box.add_run()._r.append(figure_group(doc, items, caption, total_w))


def add_figure_placeholder(doc, caption, path):
    """no picture: a one-line note for the author and the caption (the picture is inserted by hand)"""
    n = int(re.match(r'Fig\. (\d)', caption).group(1)) if re.match(r'Fig\. (\d)', caption) else 0
    names = {1: 'Fig1_mechanism', 2: 'Fig2_plant', 3: 'Fig3_control', 4: 'Fig4_pareto', 5: 'Fig5_dry', 6: 'Fig6_slip', 7: 'Fig7_drews'}.get(n, '?') + '.png'
    note = add_text(doc, 'JSTT-figure-caption', '[Insert figure here: %s]' % names, size=9, italic=True)
    note.paragraph_format.keep_with_next = True
    add_text(doc, 'JSTT-figure-caption', caption, size=9, bold=True)


def add_table(doc, caption, rows):
    cap = add_text(doc, 'JSTT-table-caption', caption, size=9, bold=True)
    cap.paragraph_format.keep_with_next = True
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol); t.alignment = 1
    plain = lambda x: re.sub(r'\$(.+?)\$', lambda q: 'x' * max(2, len(q.group(1)) // 3), x)
    col = lambda ci: [plain(r[ci]) for r in rows if ci < len(r)]
    dense = FLOW and ncol >= 5
    cmw = (lambda n: 0.135 * n + 0.2) if dense else (lambda n: 0.17 * n + 0.45)
    lo = [cmw(max(len(w) for c in col(ci) for w in (c.split() or ['']))) for ci in range(ncol)]
    hi = [max(lo[ci], cmw(max(len(c) for c in col(ci)))) for ci in range(ncol)]
    tw = COL_W if FLOW else TEXT_W
    if sum(hi) <= tw:
        widths = hi
    else:
        spare = max(0.0, tw - sum(lo)); extra = [h - l for h, l in zip(hi, lo)]
        widths = [l + spare * e / sum(extra) for l, e in zip(lo, extra)]
    wordy = lambda c: any(len(w) >= 4 and w.isalpha() and not w.isupper() for w in re.split(r'[^A-Za-z]+', re.sub(r'\$.+?\$', '', c)))
    centred = [sum(wordy(r[ci]) for r in rows[1:] if ci < len(r)) * 2 < max(1, len(rows) - 1) for ci in range(ncol)]
    tblPr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side, val in (('top', 'single'), ('left', 'nil'), ('bottom', 'single'), ('right', 'nil'), ('insideH', 'nil'), ('insideV', 'nil')):
        el = OxmlElement('w:' + side); el.set(qn('w:val'), val)
        if val == 'single':
            el.set(qn('w:sz'), '8'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblPr.append(borders)
    for ri, rw in enumerate(rows):
        trPr = t.rows[ri]._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
        for ci in range(ncol):
            cell = t.rows[ri].cells[ci]; cell.width = Cm(widths[ci]); p = cell.paragraphs[0]
            p.style = doc.styles['JSTT-table-col-head' if ri == 0 else 'JSTT-table-text']
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or centred[ci]) else WD_ALIGN_PARAGRAPH.LEFT
            add_rich(p, rw[ci] if ci < len(rw) else '', bold=(ri == 0))
            set_size(p, 7.5 if dense else 9)
            if ri < len(rows) - 1:
                p.paragraph_format.keep_with_next = True
            if ri == 0:                                   # rule under the header row
                tcPr = cell._tc.get_or_add_tcPr(); tb = OxmlElement('w:tcBorders'); b = OxmlElement('w:bottom')
                b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:space'), '0'); b.set(qn('w:color'), 'auto'); tb.append(b); tcPr.append(tb)
    sp = doc.add_paragraph(style='10.JSTT-Text'); sp.paragraph_format.space_before = Pt(0); set_size(sp, 4)


def add_reference(doc, text):
    p = doc.add_paragraph(style='EndNote Bibliography')
    add_rich(p, text)
    p.paragraph_format.left_indent = Cm(0.75); p.paragraph_format.first_line_indent = Cm(-0.75); p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        r.font.size = Pt(9); r.font.name = 'Arial'


# ------------------------------------------------------------------------------------------------ parsing
def parse(lines):
    sep = re.compile(r'^\s*(=|-){20,}\s*$')
    h1 = re.compile(r'^(\d+)\.\s+(\S.*)$')
    h2 = re.compile(r'^(\d+)\.(\d+)\.\s+(\S.*)$')
    eq = re.compile(r'^\s*\((\d+)\)\s+(.*)$')
    latex = re.compile(r'^\s*\[\s(.*)\s\]\s*$')
    figcap = re.compile(r'^\s*\[((?:Fig\.|Figure)\s+\d+[.:].*)\]\s*$')
    tabcap = re.compile(r'^\s*(Table\s+\d+[.:].*)$')
    seps = [i for i, l in enumerate(lines) if re.match(r'^\s*={20,}\s*$', l)]
    title = lines[seps[0] + 1].strip()
    blocks, front = [], {'title': title}
    i = seps[1] + 1
    section = None
    while i < len(lines):
        raw = lines[i]; s = raw.strip()
        if not s or sep.match(raw):
            i += 1; continue
        if s == 'Abstract':
            section = 'abstract'; i += 1; continue
        if section == 'abstract':
            if s.startswith('Keywords:'):
                front['keywords'] = s[len('Keywords:'):].strip().rstrip('.')
                section = None
            else:
                front['abstract'] = s
            i += 1; continue
        if s in UNNUMBERED_H:
            blocks.append({'kind': 'uh', 'text': s}); i += 1; continue
        if h2.match(s) and len(s) < 120:
            blocks.append({'kind': 'h2', 'text': h2.match(s).group(3)}); i += 1; continue
        if h1.match(s) and len(s) < 100 and not re.match(r'^\d+\.\d', s):
            blocks.append({'kind': 'h1', 'text': h1.match(s).group(2)}); i += 1; continue
        me = eq.match(raw)
        if me and i + 1 < len(lines) and latex.match(lines[i + 1]):
            n = int(me.group(1))
            blocks.append({'kind': 'eq', 'n': n, 'latex': latex.match(lines[i + 1]).group(1), 'wide': n in WIDE_EQ}); i += 2; continue
        mf = figcap.match(raw)
        if mf:
            cap, _, path = mf.group(1).partition(' | ')
            blocks.append({'kind': 'fig', 'caption': cap.strip(), 'path': path.strip(), 'wide': (not FLOW) and '_col.' not in path.split(' ; ')[0]}); i += 1; continue
        mt = tabcap.match(raw)
        if mt and i + 1 < len(lines) and ' | ' in lines[i + 1]:
            rows = []; j = i + 1
            while j < len(lines) and ' | ' in lines[j]:
                rows.append([c.strip() for c in lines[j].strip().split(' | ')]); j += 1
            blocks.append({'kind': 'tab', 'caption': mt.group(1), 'rows': rows, 'wide': not FLOW}); i = j; continue
        if s.startswith('- '):
            blocks.append({'kind': 'bullet', 'text': s[2:]}); i += 1; continue
        if re.match(r'^\[\d+\]\.', s):
            blocks.append({'kind': 'ref', 'text': s}); i += 1; continue
        if s.startswith('\\noindent '):
            blocks.append({'kind': 'p', 'text': s[len('\\noindent '):], 'indent0': True}); i += 1; continue
        blocks.append({'kind': 'p', 'text': s}); i += 1
    return front, blocks


# ------------------------------------------------------------------------------------------------ build
def build(txt_path, docx_path):
    lines = open(txt_path, encoding='utf-8').read().splitlines()
    front, blocks = parse(lines)
    doc = Document(TEMPLATE)
    body = doc.element.body
    two_col = None
    for sp in body.iter(qn('w:sectPr')):
        if sp.getparent().tag == qn('w:pPr'):
            two_col = sp; break
    import copy
    keep_sect = copy.deepcopy(two_col)
    tbl = doc.tables[0]._tbl
    for ch in list(body):
        if ch is not tbl:
            body.remove(ch)
    body.append(keep_sect)
    fill_front(doc, front['title'], front['abstract'], front['keywords'])
    add_text(doc, '10.JSTT-Text', ABBREVIATIONS, indent0=True, size=8, italic=True)
    lay = Layout(doc)
    i = 0
    while i < len(blocks):
        b = blocks[i]
        wide = b.get('wide', False)
        lay.to(1 if wide else 2)
        k = b['kind']
        if k == 'h1':
            add_text(doc, '4.JSST-1storder-head', b['text'])
        elif k == 'h2':
            add_text(doc, '4.1JSTT-2ndorder-head', b['text'])
        elif k == 'uh':
            add_text(doc, 'JSTT-acknowledgement', b['text'])
        elif k == 'p':
            add_text(doc, '10.JSTT-Text', b['text'], indent0=b.get('indent0', False))
        elif k == 'bullet':
            add_text(doc, 'JSTT-bulletlist', b['text'])
        elif k == 'eq':
            jstt_equation(doc, b['latex'], b['n'])
        elif k == 'fig':
            if FLOW:
                add_figure_placeholder(doc, b['caption'], b['path'])
            else:
                add_figure(doc, b['caption'], b['path'], wide=b['wide'])
        elif k == 'tab':
            add_table(doc, b['caption'], b['rows'])
        elif k == 'ref':
            add_reference(doc, b['text'])
        i += 1
    last = body.findall(qn('w:sectPr'))[-1]
    strip_title_pg(last) if lay.breaks else None
    doc.save(docx_path)
    patch_headers(docx_path)
    print('Wrote', docx_path)


def patch_headers(path):
    """running heads: volume year and first author instead of the template placeholders"""
    tmp = path + '.tmp'
    zin = zipfile.ZipFile(path); zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        data = zin.read(it.filename)
        if re.match(r'word/header[12]\.xml', it.filename):
            x = data.decode('utf-8')
            state = {'year': False}

            def fix(mm):                                  # the template splits the running head over many runs
                t = mm.group(2)
                if t == ' (202), ':
                    t = ' (2026), '
                elif t == ' (202':
                    t = ' (2026'; state['year'] = True
                elif t == 'X' and state['year']:
                    t = ''; state['year'] = False
                    t = ''
                elif t == ' All page':
                    t = 'page'
                elif t == 'Author':
                    t = 'Tran'
                elif t == '. et al':
                    t = ' et al.'
                elif t == '. ':
                    t = ' '
                elif t == ' al':
                    t = ' al.'
                opening = mm.group(1)
                if 'xml:space' not in opening:
                    opening = opening[:-1] + ' xml:space="preserve">'
                return opening + t + mm.group(3)
            x = re.sub(r'(<w:t(?: [^>]*)?>)([^<]*)(</w:t>)', fix, x)
            data = x.encode('utf-8')
        zout.writestr(it, data)
    zout.close(); zin.close()
    os.replace(tmp, path)


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
